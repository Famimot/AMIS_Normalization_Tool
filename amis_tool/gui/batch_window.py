"""
Batch AMIS Processing Window v6
WITH TOOLTIPS, AUTO COLUMN WIDTH AND STEP HIGHLIGHTING

=== CHANGES (v6) ===
+ raw_nohdr: no ID column, column names taken from file
+ Bold font on completed steps (4 buttons)
+ Font reset in "Clear list"

=== KEPT FROM v3.6 ===
= 5 modes: raw_single, raw_one, raw_multi, raw_nohdr, aggregated
= Detailed tooltips with file structure examples
= Table headers without markers
= "Select folder" button removed (manual selection only)
= Column widths auto-fit content
= Excel pair coloring
= click-to-graph: click header → AMIS graph

=== COMPATIBILITY ===
With amis_calculations.py v6.1: YES (raw_to_amis)
With run_amis.py: YES (BatchAMISApp)

=== REQUIREMENTS ===
pip install openpyxl   (for Excel coloring)
"""

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import os
import sys
import pandas as pd
import numpy as np
from datetime import datetime
from scipy.interpolate import interp1d
import matplotlib
matplotlib.use('TkAgg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure

try:
    from amis_tool.gui.dialogs import DataFormatDialog
    from amis_tool.core.amis_calculations import (
        amis_conversion,
        safe_aggregated_to_raw,
        raw_to_amis,
    )
    from amis_tool.utils.helpers import center_window
    AMIS_AVAILABLE = True
except ImportError:
    AMIS_AVAILABLE = False
    print("Warning: AMIS modules not available")


# =========================================================
# TIME PARSER
# =========================================================

def parse_time_to_seconds(val):
    """'H:MM:SS' or 'MM:SS' → seconds."""
    if val is None:
        return np.nan
    if isinstance(val, (int, float, np.integer, np.floating)):
        return float(val)
    s = str(val).strip()
    if not s:
        return np.nan
    if ':' in s:
        parts = s.split(':')
        try:
            if len(parts) == 3:
                h, m, sec = parts
                return int(h) * 3600 + int(m) * 60 + float(sec)
            elif len(parts) == 2:
                m, sec = parts
                return int(m) * 60 + float(sec)
        except Exception:
            return np.nan
    try:
        return float(s)
    except Exception:
        return np.nan


def column_to_numeric(series):
    """Column → numbers; if ':' present — parse as time."""
    sample = series.dropna().astype(str).head(50)
    looks_like_time = sample.str.contains(':', regex=False).any()
    if looks_like_time:
        parsed = series.apply(parse_time_to_seconds)
        return pd.to_numeric(parsed, errors='coerce').dropna().values
    else:
        return pd.to_numeric(series, errors='coerce').dropna().values


class ToolTip:
    """Create tooltip for widgets"""
    def __init__(self, widget, text):
        self.widget = widget
        self.text = text
        self.tip_window = None
        widget.bind('<Enter>', self.enter)
        widget.bind('<Leave>', self.leave)

    def enter(self, event=None):
        x, y, _, _ = self.widget.bbox("insert")
        x += self.widget.winfo_rootx() + 25
        y += self.widget.winfo_rooty() + 25
        self.tip_window = tw = tk.Toplevel(self.widget)
        tw.wm_overrideredirect(True)
        tw.wm_geometry(f"+{x}+{y}")
        label = tk.Label(tw, text=self.text, justify=tk.LEFT,
                         background="#ffffe0", relief=tk.SOLID, borderwidth=1,
                         font=("Arial", 9), padx=6, pady=4)
        label.pack()

    def leave(self, event=None):
        if self.tip_window:
            self.tip_window.destroy()
            self.tip_window = None


class BatchFileConfigDialog(tk.Toplevel):
    """Single file settings dialog"""

    def __init__(self, parent, file_path, current_settings=None):
        super().__init__(parent)
        self.title(f"Settings: {os.path.basename(file_path)}")
        self.geometry("560x850")
        self.file_path = file_path
        self.result = None

        try:
            if file_path.lower().endswith('.csv'):
                self.df = pd.read_csv(file_path, encoding='utf-8')
            else:
                self.df = pd.read_excel(file_path)
        except Exception as e:
            messagebox.showerror("Error", f"Cannot load file:\n{str(e)}")
            self.df = pd.DataFrame()

        self.n_cols = len(self.df.columns)

        self.data_type_var = tk.StringVar(value="raw_multi")
        self.scale_mode_var = tk.StringVar(value="auto")
        self.method_var = tk.StringVar(value="mean")

        if self.n_cols == 1:
            self.data_type_var.set("raw_single")
        elif self.n_cols == 2:
            self.data_type_var.set("raw_one")

        self.create_widgets()
        center_window(self)
        self.transient(parent)
        self.grab_set()

        self.after(100, lambda: self.ok_button.focus_set())

    def create_widgets(self):
        main_frame = ttk.Frame(self, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(main_frame,
                 text=f"File: {os.path.basename(self.file_path)}",
                 font=('Arial', 12, 'bold')).pack(anchor=tk.W, pady=(0, 5))

        file_info = f"{len(self.df)} rows, {self.n_cols} columns"
        ttk.Label(main_frame, text=file_info,
                 font=('Arial', 9)).pack(anchor=tk.W, pady=(0, 15))

        preview_frame = ttk.LabelFrame(main_frame, text="Data preview", padding="8")
        preview_frame.pack(fill=tk.X, pady=(0, 15))

        preview_text = tk.Text(preview_frame, height=6, wrap=tk.NONE,
                             font=('Consolas', 8))
        preview_scrollx = ttk.Scrollbar(preview_frame, orient=tk.HORIZONTAL,
                                       command=preview_text.xview)
        preview_scrolly = ttk.Scrollbar(preview_frame, orient=tk.VERTICAL,
                                       command=preview_text.yview)
        preview_text.configure(xscrollcommand=preview_scrollx.set,
                             yscrollcommand=preview_scrolly.set)

        preview_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        preview_scrolly.pack(side=tk.RIGHT, fill=tk.Y)
        preview_scrollx.pack(side=tk.BOTTOM, fill=tk.X)

        preview_text.insert(tk.END, self.df.head().to_string())
        preview_text.config(state=tk.DISABLED)

        ttk.Separator(main_frame, orient=tk.HORIZONTAL).pack(fill=tk.X, pady=10)

        if self.n_cols >= 2:
            values = column_to_numeric(self.df.iloc[:, 1])
        elif self.n_cols == 1:
            values = column_to_numeric(self.df.iloc[:, 0])
        else:
            values = np.array([])

        data_min = float(np.min(values)) if len(values) > 0 else 0
        data_max = float(np.max(values)) if len(values) > 0 else 1

        ttk.Label(main_frame, text="Normalization settings:",
                 font=('Arial', 10, 'bold')).pack(anchor=tk.W, pady=(0, 10))

        # ============ 1. DATA FORMAT ============
        format_frame = ttk.LabelFrame(main_frame, text="1. Data Format", padding="8")
        format_frame.pack(fill=tk.X, pady=(0, 10))

        rb_single = ttk.Radiobutton(
            format_frame,
            text="● Raw: ONE column only (no ID)",
            variable=self.data_type_var, value="raw_single")
        rb_single.pack(anchor=tk.W, pady=2)
        ToolTip(rb_single,
            "File contains EXACTLY ONE data column.\n"
            "No ID is created — output will have 2 columns:\n"
            "   <column name> | AMIS <column name>\n\n"
            "Example file:\n"
            "  Time\n"
            "  0:00:23\n"
            "  0:00:45\n"
            "  0:01:02")

        rb_raw1 = ttk.Radiobutton(
            format_frame,
            text="○ Raw: ID + ONE value column",
            variable=self.data_type_var, value="raw_one")
        rb_raw1.pack(anchor=tk.W, pady=2)
        ToolTip(rb_raw1,
            "File contains EXACTLY TWO columns:\n"
            "  Column 1 — ID/names (not processed)\n"
            "  Column 2 — values for normalization\n\n"
            "Output: ID | <column> | AMIS <column>\n\n"
            "Example file:\n"
            "  ID    Time\n"
            "  1     0:00:23\n"
            "  2     0:00:45\n"
            "  3     0:01:02")

        rb_rawN = ttk.Radiobutton(
            format_frame,
            text="○ Raw: ID + MULTIPLE value columns",
            variable=self.data_type_var, value="raw_multi")
        rb_rawN.pack(anchor=tk.W, pady=2)
        ToolTip(rb_rawN,
            "File contains SEVERAL columns:\n"
            "  Column 1 — ID (kept as ID)\n"
            "  Columns 2..N — data (each = separate series)\n\n"
            "Output: ID | P1 | AMIS P1 | P2 | AMIS P2 | ...\n\n"
            "Example file:\n"
            "  ID    Temperature   Pressure\n"
            "  1     23.5          101.2\n"
            "  2     24.1          101.5\n"
            "  3     22.8          100.9")

        rb_rawNoHdr = ttk.Radiobutton(
            format_frame,
            text="○ Raw: MANY columns WITHOUT ID",
            variable=self.data_type_var, value="raw_nohdr")
        rb_rawNoHdr.pack(anchor=tk.W, pady=2)
        ToolTip(rb_rawNoHdr,
            "No ID column — ALL columns are data.\n"
            "Column names are taken from file headers.\n\n"
            "Output: P1 | AMIS P1 | P2 | AMIS P2 | ...\n"
            "(no ID column)\n\n"
            "Example file:\n"
            "  Temperature   Pressure\n"
            "  23.5          101.2\n"
            "  24.1          101.5\n"
            "  22.8          100.9")

        if self.n_cols == 1:
            rb_agg = ttk.Radiobutton(
                format_frame,
                text="○ Aggregated (NOT AVAILABLE for 1 column)",
                variable=self.data_type_var, value="aggregated",
                state=tk.DISABLED)
            rb_agg.pack(anchor=tk.W, pady=2)
        else:
            rb_agg = ttk.Radiobutton(
                format_frame,
                text="○ Aggregated: values + frequencies",
                variable=self.data_type_var, value="aggregated")
            rb_agg.pack(anchor=tk.W, pady=2)
            ToolTip(rb_agg,
                "Data is already grouped:\n"
                "  Column 1 — values\n"
                "  Column 2 — frequency (how many times each occurs)\n\n"
                "Values are expanded into a full series on output.\n\n"
                "Example file:\n"
                "  Value     Frequency\n"
                "  10.5      3\n"
                "  11.2      7\n"
                "  12.8      2")

        if self.n_cols == 1:
            recommend = "💡 Detected 1 column → recommended ● raw_single"
        elif self.n_cols == 2:
            recommend = "💡 Detected 2 columns → recommended ○ raw_one"
        else:
            recommend = f"💡 Detected {self.n_cols} columns → recommended ○ raw_multi"
        ttk.Label(format_frame, text=recommend,
                 font=('Arial', 8), foreground="#2c3e50").pack(anchor=tk.W, pady=(8, 0))

        # ============ 2. SCALE BOUNDS ============
        scale_frame = ttk.LabelFrame(main_frame, text="2. Scale Bounds", padding="8")
        scale_frame.pack(fill=tk.X, pady=(0, 10))

        data_info = f"Data range: from {data_min:.4g} to {data_max:.4g}"
        ttk.Label(scale_frame, text=data_info,
                 font=('Arial', 9)).pack(anchor=tk.W, pady=(0, 10))

        rb_auto = ttk.Radiobutton(scale_frame, text="Auto: use data min/max",
                       variable=self.scale_mode_var, value="auto",
                       command=self.toggle_scale_mode)
        rb_auto.pack(anchor=tk.W, pady=2)
        ToolTip(rb_auto,
            "0 AMIS = minimum value in data\n"
            "100 AMIS = maximum value in data")

        rb_manual = ttk.Radiobutton(scale_frame,
                       text="Manual: specify values for 0 and 100 AMIS",
                       variable=self.scale_mode_var, value="manual",
                       command=self.toggle_scale_mode)
        rb_manual.pack(anchor=tk.W, pady=2)
        ToolTip(rb_manual,
            "Manually set which values correspond to\n"
            "0 AMIS and 100 AMIS (for reference comparison)")

        self.manual_frame = ttk.Frame(scale_frame)
        self.manual_frame.pack(fill=tk.X, pady=(10, 0))

        ttk.Label(self.manual_frame, text="0 AMIS =").grid(row=0, column=0, padx=(0, 5), sticky=tk.W)
        self.min_entry = ttk.Entry(self.manual_frame, width=12, state=tk.DISABLED)
        self.min_entry.grid(row=0, column=1, padx=5, pady=2)

        ttk.Label(self.manual_frame, text="100 AMIS =").grid(row=0, column=2, padx=(15, 5), sticky=tk.W)
        self.max_entry = ttk.Entry(self.manual_frame, width=12, state=tk.DISABLED)
        self.max_entry.grid(row=0, column=3, padx=5, pady=2)

        quick_btn = ttk.Button(scale_frame, text="Use file min/max",
                  command=lambda: self.fill_with_data_bounds(data_min, data_max),
                  width=35)
        quick_btn.pack(pady=2)
        ToolTip(quick_btn, "Fill fields with file minimum and maximum")

        # ============ 3. METHOD ============
        method_frame = ttk.LabelFrame(main_frame, text="3. Normalization Method", padding="8")
        method_frame.pack(fill=tk.X, pady=(0, 15))

        rb_mean = ttk.Radiobutton(method_frame, text="Mean (classical AMIS)",
                                 variable=self.method_var, value="mean")
        rb_mean.pack(anchor=tk.W, pady=2)
        ToolTip(rb_mean, "Classical AMIS: uses arithmetic mean")

        info_text = "💡 AMIS models: 3-point (6+), 5-point (10+), 9-point (20+), 17-point (40+)"
        ttk.Label(method_frame, text=info_text,
                 font=('Arial', 8), foreground="gray").pack(anchor=tk.W, pady=(10, 0))

        # BUTTONS
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=(20, 0))

        center_frame = ttk.Frame(button_frame)
        center_frame.pack(expand=True)

        self.ok_button = ttk.Button(center_frame, text="Save", width=12, command=self.ok)
        self.ok_button.pack(side=tk.LEFT, padx=5)

        cancel_btn = ttk.Button(center_frame, text="Cancel", width=12, command=self.cancel)
        cancel_btn.pack(side=tk.LEFT, padx=5)

        help_btn = ttk.Button(center_frame, text="Help", width=12, command=self.show_help)
        help_btn.pack(side=tk.LEFT, padx=5)

    def toggle_scale_mode(self):
        if self.scale_mode_var.get() == "manual":
            self.min_entry.config(state=tk.NORMAL)
            self.max_entry.config(state=tk.NORMAL)
        else:
            self.min_entry.config(state=tk.DISABLED)
            self.max_entry.config(state=tk.DISABLED)

    def fill_with_data_bounds(self, data_min, data_max):
        self.min_entry.delete(0, tk.END)
        self.max_entry.delete(0, tk.END)
        self.min_entry.insert(0, f"{data_min:.4g}")
        self.max_entry.insert(0, f"{data_max:.4g}")

    def show_help(self):
        help_text = """📖 FILE SETTINGS (v6):

DATA FORMAT — 5 MODES:

● raw_single — 1 column
  File: values only, no ID.
  Output: <column> | AMIS <column>

○ raw_one — ID + 1 column
  File: ID | values
  Output: ID | <column> | AMIS <column>

○ raw_multi — ID + multiple columns
  File: ID | P1 | P2 | ...
  Output: ID | P1 | AMIS P1 | P2 | AMIS P2 | ...

○ raw_nohdr — multiple columns, no ID
  File: P1 | P2 | ...  (headers are kept)
  Output: P1 | AMIS P1 | P2 | AMIS P2 | ...  (no ID)

○ aggregated — values + frequencies
  File: value | frequency
  Output: expanded into full series

GRAPH:
  Click on header → AMIS graph for that parameter.

TIME COLUMNS:
  Values like 0:00:23 are automatically converted to seconds.

AMIS MODELS:
  3-point: 6+, 5-point: 10+, 9-point: 20+, 17-point: 40+"""
        messagebox.showinfo("Help - File Settings", help_text)

    def ok(self):
        data_type = self.data_type_var.get()
        method = self.method_var.get()

        if data_type == "raw_single" and self.n_cols != 1:
            if not messagebox.askyesno("Warning",
                f"Mode raw_single is intended for files with 1 column,\n"
                f"but file has {self.n_cols} columns.\n\n"
                f"Continue? (first column will be used)"):
                return
        if data_type == "raw_one" and self.n_cols != 2:
            if not messagebox.askyesno("Warning",
                f"Mode raw_one is intended for files with 2 columns,\n"
                f"but file has {self.n_cols} columns.\n\n"
                f"Continue?"):
                return

        if self.scale_mode_var.get() == "auto":
            scale_min = None
            scale_max = None
        else:
            try:
                scale_min = float(self.min_entry.get()) if self.min_entry.get().strip() else None
                scale_max = float(self.max_entry.get()) if self.max_entry.get().strip() else None
                if scale_min is None or scale_max is None:
                    messagebox.showerror("Error", "Fill both fields for manual mode")
                    return
                if scale_min >= scale_max:
                    messagebox.showerror("Error", "Minimum must be LESS than maximum")
                    return
            except ValueError:
                messagebox.showerror("Error", "Enter valid numbers")
                return

        self.result = {
            'file_path': self.file_path,
            'data_type': data_type,
            'scale_mode': self.scale_mode_var.get(),
            'scale_min': scale_min,
            'scale_max': scale_max,
            'method': method
        }
        self.destroy()

    def cancel(self):
        self.result = None
        self.destroy()

    def get_result(self):
        self.wait_window(self)
        return self.result


class BatchAMISApp(tk.Tk):
    """Batch processing main window (v.6 — no-hdr without ID, step highlighting)"""

    def __init__(self):
        super().__init__()

        self.title("AMIS - Batch Processing v6")
        self.geometry("1400x900")
        self.minsize(1200, 750)

        self.input_folder = None
        self.file_paths = []
        self.file_configs = {}
        self.file_data = {}
        self.results = {}
        self.combined_df = None
        self.id_column_name = "ID"
        self.has_id_column = False
        self.column_pairs = []

        self.default_settings = {
            'data_type': 'raw_multi',
            'scale_mode': 'auto',
            'scale_min': None,
            'scale_max': None,
            'method': 'mean'
        }

        self.figure = None
        self.ax = None
        self.canvas = None

        self._init_styles()

        self.create_widgets()
        center_window(self)

        self.protocol("WM_DELETE_WINDOW", self.on_close)

    def _init_styles(self):
        """Button styles: normal and completed (bold)."""
        style = ttk.Style()
        base_font = ("Arial", 9, "normal")
        done_font = ("Arial", 9, "bold")
        style.configure("StepPending.TButton", font=base_font)
        style.configure("StepDone.TButton", font=done_font)

    def _set_step_done(self, btn, done=True):
        """Toggle button between normal and bold font."""
        btn.configure(style="StepDone.TButton" if done else "StepPending.TButton")

    def create_widgets(self):
        main_paned = ttk.PanedWindow(self, orient=tk.HORIZONTAL)
        main_paned.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        left_panel = ttk.Frame(main_paned, width=280)
        main_paned.add(left_panel, weight=1)

        right_panel = ttk.Frame(main_paned)
        main_paned.add(right_panel, weight=3)

        # LEFT
        control_frame = ttk.Frame(left_panel, padding="8")
        control_frame.pack(fill=tk.BOTH, expand=True)

        ttk.Label(control_frame, text="BATCH PROCESSING v6",
                 font=('Arial', 10, 'bold')).pack(pady=(0, 8))

        ttk.Label(control_frame, text="Column-wise AMIS\n(click header → graph)",
                 font=('Arial', 8), foreground="blue", justify=tk.CENTER).pack(pady=(0, 10))

        file_frame = ttk.LabelFrame(control_frame, text="📁 Files", padding="5")
        file_frame.pack(fill=tk.X, pady=(0, 6))

        self.folder_label = ttk.Label(file_frame, text="No files selected",
                                     font=('Arial', 8), foreground="gray")
        self.folder_label.pack(anchor=tk.W, pady=(0, 5))

        self.btn_select = ttk.Button(file_frame, text="Select files manually",
                                     command=self.select_files_manually,
                                     style="StepPending.TButton")
        self.btn_select.pack(fill=tk.X, pady=2)
        ToolTip(self.btn_select, "Select one or more files (xlsx / xls / csv)")

        self.btn_clear = ttk.Button(file_frame, text="Clear list",
                                    command=self.clear_file_list)
        self.btn_clear.pack(fill=tk.X, pady=2)
        ToolTip(self.btn_clear, "Clear file list and reset results")

        ttk.Separator(file_frame, orient='horizontal').pack(fill=tk.X, pady=5)

        self.btn_config = ttk.Button(file_frame, text="⚙️ Configure ALL files",
                                     command=self.configure_all_files,
                                     style="StepPending.TButton")
        self.btn_config.pack(fill=tk.X, pady=2)
        ToolTip(self.btn_config, "Configure data format, bounds and method for each file")

        status_frame = ttk.LabelFrame(control_frame, text="📋 File Status", padding="5")
        status_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 8))

        self.status_tree = ttk.Treeview(status_frame, height=12, show='tree')
        status_scroll = ttk.Scrollbar(status_frame, orient=tk.VERTICAL,
                                     command=self.status_tree.yview)
        self.status_tree.configure(yscrollcommand=status_scroll.set)

        self.status_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        status_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        self.status_tree['columns'] = ('status', 'type', 'bounds', 'method')
        self.status_tree.column('#0', width=150, minwidth=120)
        self.status_tree.column('status', width=40, minwidth=30)
        self.status_tree.column('type', width=70, minwidth=60)
        self.status_tree.column('bounds', width=80, minwidth=70)
        self.status_tree.column('method', width=60, minwidth=50)

        self.status_tree.heading('#0', text='File')
        self.status_tree.heading('status', text='✅')
        self.status_tree.heading('type', text='Type')
        self.status_tree.heading('bounds', text='Bounds')
        self.status_tree.heading('method', text='Method')

        process_frame = ttk.LabelFrame(control_frame, text="🔄 Processing", padding="5")
        process_frame.pack(fill=tk.X, pady=(0, 6))

        self.btn_load = ttk.Button(process_frame, text="Load and analyze",
                                   command=self.load_and_analyze,
                                   style="StepPending.TButton")
        self.btn_load.pack(fill=tk.X, pady=2)
        ToolTip(self.btn_load, "Load all selected files and parse their columns")

        self.btn_norm = ttk.Button(process_frame, text="Normalize AMIS",
                                   command=self.normalize_files,
                                   style="StepPending.TButton")
        self.btn_norm.pack(fill=tk.X, pady=2)
        ToolTip(self.btn_norm, "Normalize all loaded series to AMIS scale 0–100")

        export_frame = ttk.LabelFrame(control_frame, text="💾 Export", padding="5")
        export_frame.pack(fill=tk.X, pady=(0, 8))

        btn_save = ttk.Button(export_frame, text="Save Excel (with colors)",
                              command=self.save_excel)
        btn_save.pack(fill=tk.X, pady=2)
        ToolTip(btn_save, "Save combined table to Excel\n(pair columns will be colored)")

        btn_copy = ttk.Button(export_frame, text="Copy table",
                              command=self.copy_to_clipboard)
        btn_copy.pack(fill=tk.X, pady=2)
        ToolTip(btn_copy, "Copy combined table to clipboard")

        ttk.Button(control_frame, text="🏠 Main menu",
                   command=self.return_to_main).pack(fill=tk.X, pady=8)

        # RIGHT
        right_paned = ttk.PanedWindow(right_panel, orient=tk.VERTICAL)
        right_paned.pack(fill=tk.BOTH, expand=True)

        table_frame = ttk.LabelFrame(right_paned,
            text="Combined table (click header → graph)",
            padding="6")
        right_paned.add(table_frame, weight=4)

        table_container = ttk.Frame(table_frame)
        table_container.pack(fill=tk.BOTH, expand=True)

        self.tree = ttk.Treeview(table_container, show='headings')

        v_scroll = ttk.Scrollbar(table_container, orient=tk.VERTICAL, command=self.tree.yview)
        h_scroll = ttk.Scrollbar(table_container, orient=tk.HORIZONTAL, command=self.tree.xview)
        self.tree.configure(yscrollcommand=v_scroll.set, xscrollcommand=h_scroll.set)

        self.tree.grid(row=0, column=0, sticky='nsew')
        v_scroll.grid(row=0, column=1, sticky='ns')
        h_scroll.grid(row=1, column=0, sticky='ew', columnspan=2)

        table_container.grid_rowconfigure(0, weight=1)
        table_container.grid_columnconfigure(0, weight=1)

        self.tree.bind('<Button-1>', self.on_tree_click)

        # INFO PANEL
        info_frame = ttk.LabelFrame(right_paned, text="ℹ️ Hint", padding="6")
        right_paned.add(info_frame, weight=1)

        hint_text = (
            "• Click ANY header → AMIS graph for that parameter.\n"
            "• When saving to Excel, pairs are automatically colored (blue / white).\n"
            "• Column widths auto-fit content; scrollbar for scrolling.\n"
            "• Completed steps are highlighted in bold."
        )
        ttk.Label(info_frame, text=hint_text,
                 font=('Arial', 9), justify=tk.LEFT,
                 foreground="#2c3e50").pack(anchor=tk.W, padx=5, pady=5)

        # LOG
        log_frame = ttk.LabelFrame(self, text="📝 Processing Log", padding="6")
        log_frame.pack(fill=tk.X, padx=5, pady=(0, 5))

        self.log_text = tk.Text(log_frame, height=6, wrap=tk.WORD,
                               font=('Consolas', 8))
        log_scroll = ttk.Scrollbar(log_frame, orient=tk.VERTICAL,
                                  command=self.log_text.yview)
        self.log_text.configure(yscrollcommand=log_scroll.set)

        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        log_scroll.pack(side=tk.RIGHT, fill=tk.Y)

        self.log_text.tag_config("error", foreground="red")
        self.log_text.tag_config("warning", foreground="orange")
        self.log_text.tag_config("success", foreground="green")
        self.log_text.tag_config("info", foreground="blue")

        self.status_var = tk.StringVar(value="Ready to work")
        status_bar = ttk.Label(self, textvariable=self.status_var,
                              relief=tk.SUNKEN, font=('Arial', 8))
        status_bar.pack(side=tk.BOTTOM, fill=tk.X, padx=5, pady=2)

    # =========================================================
    # Utilities
    # =========================================================

    def on_close(self):
        if messagebox.askyesno("Confirm",
                              "Return to main menu?\nUnsaved data will be lost."):
            self.return_to_main()
        else:
            self.destroy()

    def log(self, message, level="info"):
        timestamp = datetime.now().strftime("%H:%M:%S")
        if level == "error":
            prefix = f"[{timestamp}] ❌ "; tag = "error"
        elif level == "warning":
            prefix = f"[{timestamp}] ⚠️ "; tag = "warning"
        elif level == "success":
            prefix = f"[{timestamp}] ✅ "; tag = "success"
        else:
            prefix = f"[{timestamp}] ℹ️ "; tag = "info"
        self.log_text.insert(tk.END, prefix + message + "\n", tag)
        self.log_text.see(tk.END)
        self.status_var.set(message)

    def select_files_manually(self):
        files = filedialog.askopenfilenames(
            title="Select files for processing",
            filetypes=[("Excel/CSV", "*.xlsx *.xls *.csv"), ("All files", "*.*")]
        )
        if files:
            self.file_paths = list(files)
            self.input_folder = os.path.dirname(files[0])
            self.folder_label.config(text=f"Selected: {len(files)} files")
            self.log(f"Files selected: {len(files)}", "success")
            self.update_file_status_list()
            self._set_step_done(self.btn_select, True)

    def update_file_status_list(self):
        for item in self.status_tree.get_children():
            self.status_tree.delete(item)
        if not self.file_paths:
            return
        for path in self.file_paths:
            filename = os.path.basename(path)
            if filename in self.file_configs:
                config = self.file_configs[filename]
                status = "✅"
                data_type = {
                    'raw_single': 'raw1col',
                    'raw_one': 'raw2col',
                    'raw_multi': 'rawN',
                    'raw_nohdr': 'nohdr',
                    'aggregated': 'agg.'
                }.get(config['data_type'], config['data_type'])
                bounds = "auto" if config['scale_mode'] == 'auto' else "man."
                method = "mean"
            else:
                status = "❓"; data_type = "not set"; bounds = "not set"; method = "?"
            self.status_tree.insert('', 'end', iid=filename, text=filename,
                                   values=(status, data_type, bounds, method))

    def configure_all_files(self):
        if not self.file_paths:
            messagebox.showwarning("No files", "Select files first")
            return
        self.log("Starting file configuration...")
        configured = 0
        for path in self.file_paths:
            filename = os.path.basename(path)
            dialog = BatchFileConfigDialog(self, path)
            result = dialog.get_result()
            if result:
                self.file_configs[filename] = {
                    'data_type': result['data_type'],
                    'scale_mode': result['scale_mode'],
                    'scale_min': result['scale_min'],
                    'scale_max': result['scale_max'],
                    'method': 'mean'
                }
                configured += 1
                self.log(f"Configured: {filename} ({result['data_type']})")
            else:
                self.log(f"Skipped: {filename} (cancelled)", "warning")
        self.update_file_status_list()
        self.log(f"✅ Files configured: {configured}/{len(self.file_paths)}", "success")
        self._set_step_done(self.btn_config, True)

    def clear_file_list(self):
        self.file_paths = []
        self.file_configs = {}
        self.file_data = {}
        self.results = {}
        self.combined_df = None
        self.column_pairs = []
        self.has_id_column = False
        self.folder_label.config(text="No files selected")
        self.update_file_status_list()
        self.update_table_display()
        self.log("File list cleared")
        self._set_step_done(self.btn_select, False)
        self._set_step_done(self.btn_config, False)
        self._set_step_done(self.btn_load, False)
        self._set_step_done(self.btn_norm, False)

    # =========================================================
    # LOAD
    # =========================================================

    def load_and_analyze(self):
        if not self.file_paths:
            messagebox.showwarning("No files", "Select files first")
            return

        not_configured = [os.path.basename(p) for p in self.file_paths
                          if os.path.basename(p) not in self.file_configs]
        if not_configured:
            response = messagebox.askyesno(
                "Not all files configured",
                f"{len(not_configured)} files not configured.\n"
                f"Use default settings? (Raw: ID + MULTIPLE columns, auto, Mean)"
            )
            if not response:
                self.log("Loading cancelled", "warning")
                return
            for filename in not_configured:
                self.file_configs[filename] = self.default_settings.copy()
            self._set_step_done(self.btn_config, True)

        self.log(f"Loading {len(self.file_paths)} files...")
        self.status_var.set("Loading files...")

        self.file_data.clear()
        self.column_pairs = []
        successful = 0

        for path in self.file_paths:
            try:
                filename = os.path.basename(path)
                config = self.file_configs.get(filename, self.default_settings)

                if path.lower().endswith('.csv'):
                    df = pd.read_csv(path, encoding='utf-8')
                else:
                    df = pd.read_excel(path)

                if df.empty:
                    self.log(f"Skipping: {filename} - file empty", "warning")
                    continue

                data_type = config['data_type']
                series_list = []
                id_values = None
                id_name = None

                if data_type == 'raw_single':
                    if len(df.columns) < 1:
                        self.log(f"raw_single: no columns in {filename}", "warning")
                        continue
                    vals = column_to_numeric(df.iloc[:, 0])
                    name = str(df.columns[0]).strip() or "Values"
                    if len(vals) >= 6:
                        series_list.append((name, vals, 0))
                    else:
                        self.log(f"  col '{name}': only {len(vals)} valid (<6), skipped", "warning")

                elif data_type == 'raw_one':
                    if len(df.columns) < 2:
                        self.log(f"raw_one: file has {len(df.columns)} columns, needs 2+", "warning")
                        continue
                    id_values = df.iloc[:, 0].tolist()
                    id_name = str(df.columns[0]) if df.columns[0] else "ID"
                    vals = column_to_numeric(df.iloc[:, 1])
                    name = str(df.columns[1]).strip() or "Values"
                    if len(vals) >= 6:
                        series_list.append((name, vals, 1))
                    else:
                        self.log(f"  col '{name}': only {len(vals)} valid (<6), skipped", "warning")

                elif data_type == 'raw_multi':
                    if len(df.columns) < 2:
                        self.log(f"raw_multi: file has {len(df.columns)} columns, needs 2+", "warning")
                        continue
                    id_values = df.iloc[:, 0].tolist()
                    id_name = str(df.columns[0]) if df.columns[0] else "ID"
                    for pos, col in enumerate(df.columns[1:], start=1):
                        vals = column_to_numeric(df[col])
                        name = str(col).strip()
                        if not name or name.lower().startswith('unnamed'):
                            name = f"Column{pos}"
                        self.log(f"  col '{name}': valid={len(vals)}")
                        if len(vals) >= 6:
                            series_list.append((name, vals, pos))

                elif data_type == 'raw_nohdr':
                    id_values = None
                    id_name = None
                    for pos, col in enumerate(df.columns, start=0):
                        vals = column_to_numeric(df[col])
                        name = str(col).strip()
                        if not name or name.lower().startswith('unnamed'):
                            name = f"Column{pos+1}"
                        self.log(f"  col '{name}': valid={len(vals)}")
                        if len(vals) >= 6:
                            series_list.append((name, vals, pos))

                elif data_type == 'aggregated':
                    try:
                        values, col_name = safe_aggregated_to_raw(df, "aggregated")
                        id_values = list(range(1, len(values) + 1))
                        id_name = "N"
                        series_list.append((str(col_name), values, 0))
                    except Exception as e:
                        self.log(f"Aggregated data error in {filename}: {e}", "warning")
                        continue

                for pos, (col_name, values, col_pos) in enumerate(series_list):
                    if len(values) < 6:
                        continue
                    key = f"{filename}::{col_pos}::{col_name}"
                    self.file_data[key] = {
                        'df': df,
                        'path': path,
                        'config': config,
                        'column': col_name,
                        'col_pos': col_pos,
                        'values': values,
                        'id_values': id_values,
                        'id_name': id_name,
                        'n_points': len(values),
                        'min': float(np.min(values)),
                        'max': float(np.max(values)),
                        'mean': float(np.mean(values))
                    }
                    successful += 1
                    self.log(f"Loaded: {key} - {len(values)} points")

                if filename in self.file_configs:
                    cfg = self.file_configs[filename]
                    data_type_short = {
                        'raw_single': 'raw1col',
                        'raw_one': 'raw2col',
                        'raw_multi': 'rawN',
                        'raw_nohdr': 'nohdr',
                        'aggregated': 'agg.'
                    }.get(cfg['data_type'], '?')
                    bounds = "auto" if cfg['scale_mode'] == 'auto' else "man."
                    self.status_tree.item(filename, values=("✓", data_type_short, bounds, "mean"))

            except Exception as e:
                self.log(f"Error loading {os.path.basename(path)}: {str(e)}", "error")

        self.status_var.set(f"Loaded: {successful} series")
        if successful > 0:
            self.log(f"✅ Loading completed: {successful} series", "success")
            self._set_step_done(self.btn_load, True)
        else:
            self.log("❌ Loading failed", "error")

    # =========================================================
    # NORMALIZE
    # =========================================================

    def normalize_files(self):
        if not self.file_data:
            messagebox.showwarning("No data", "Load files first")
            return
        if not AMIS_AVAILABLE:
            messagebox.showerror("Error", "AMIS modules not available")
            return

        self.log(f"Starting normalization of {len(self.file_data)} series...")
        self.status_var.set("Normalizing...")

        self.results.clear()
        successful = 0

        for key, data in self.file_data.items():
            try:
                values = data['values']
                config = data['config']

                fixed_min = config['scale_min'] if config['scale_mode'] == 'manual' else None
                fixed_max = config['scale_max'] if config['scale_mode'] == 'manual' else None

                converted, points_coords, available_models = amis_conversion(
                    values, fixed_min=fixed_min, fixed_max=fixed_max, method='mean')

                amis_for_raw = raw_to_amis(
                    values, converted, points_coords, available_models,
                    fixed_min=fixed_min, fixed_max=fixed_max)

                self.results[key] = {
                    'column': data['column'],
                    'col_pos': data['col_pos'],
                    'filename': os.path.basename(data['path']),
                    'raw_values': values,
                    'amis_values': np.asarray(amis_for_raw, dtype=float),
                    'id_values': data.get('id_values'),
                    'id_name': data.get('id_name'),
                    'config': config,
                    'method': 'Mean',
                    'n_points': len(values),
                }
                successful += 1
                self.log(f"Normalized: {key} ({len(values)} values)")

            except Exception as e:
                self.log(f"Error normalizing {key}: {str(e)}", "error")

        self.status_var.set(f"Normalization: {successful} series")

        if successful > 0:
            self.create_combined_table()
            self.log(f"✅ Normalization completed: {successful} series", "success")
            self._set_step_done(self.btn_norm, True)
            messagebox.showinfo("Done", f"AMIS table created\nSeries: {successful}")
        else:
            self.log("❌ Normalization failed", "error")

    # =========================================================
    # BUILD TABLE
    # =========================================================

    def create_combined_table(self):
        """ID (if present) | <P1> | AMIS <P1> | <P2> | AMIS <P2> | ..."""
        if not self.results:
            self.combined_df = None
            self.column_pairs = []
            self.has_id_column = False
            self.update_table_display()
            return

        by_file = {}
        for key, r in self.results.items():
            by_file.setdefault(r['filename'], []).append(r)

        all_rows = []
        ordered_cols_global = None
        id_name_global = "ID"
        has_id_any = False

        self.column_pairs = []

        for filename, items in by_file.items():
            items.sort(key=lambda x: x.get('col_pos', 0))

            seen = {}
            named = []
            for r in items:
                base = r['column']
                if base in seen:
                    seen[base] += 1
                    name = f"{base}_{seen[base]}"
                else:
                    seen[base] = 0
                    name = base
                named.append((name, r))

            if ordered_cols_global is None:
                ordered_cols_global = [n for n, _ in named]
                for r in items:
                    if r.get('id_name'):
                        id_name_global = r['id_name']
                        break

            file_has_id = any(r.get('id_values') is not None for _, r in named)
            if file_has_id:
                has_id_any = True

            lengths = [len(r['raw_values']) for _, r in named]
            n_rows = max(lengths) if lengths else 0

            for i in range(n_rows):
                row = {}
                if file_has_id:
                    id_val = None
                    for _, r in named:
                        ids = r.get('id_values')
                        if ids is not None and i < len(ids):
                            id_val = ids[i]
                            break
                    row[id_name_global] = id_val

                for name, r in named:
                    raw = r['raw_values'][i] if i < len(r['raw_values']) else None
                    amis = r['amis_values'][i] if i < len(r['amis_values']) else None
                    row[name] = raw
                    row[f'AMIS {name}'] = round(float(amis), 4) if amis is not None else None

                all_rows.append(row)

        self.combined_df = pd.DataFrame(all_rows)
        self.has_id_column = has_id_any

        if not self.combined_df.empty and ordered_cols_global:
            ordered = []
            if has_id_any:
                ordered.append(id_name_global)
            for name in ordered_cols_global:
                ordered.append(name)
                ordered.append(f'AMIS {name}')
            existing = [c for c in ordered if c in self.combined_df.columns]
            rest = [c for c in self.combined_df.columns if c not in existing]
            self.combined_df = self.combined_df[existing + rest]

        for name in ordered_cols_global or []:
            self.column_pairs.append((name, f'AMIS {name}'))

        self.id_column_name = id_name_global if has_id_any else None

        self.update_table_display()

    def update_table_display(self):
        """Headers without markers, column width auto-fits content"""
        for item in self.tree.get_children():
            self.tree.delete(item)

        if self.combined_df is None or self.combined_df.empty:
            return

        columns = list(self.combined_df.columns)
        self.tree['columns'] = columns

        for col in columns:
            max_len = len(str(col))
            try:
                for val in self.combined_df[col].head(200):
                    if pd.isna(val):
                        continue
                    s = f"{val:.2f}" if (col.startswith('AMIS ') and isinstance(val, (int, float, np.floating))) else str(val)
                    max_len = max(max_len, len(s))
            except Exception:
                pass

            width = max(80, min(max_len * 9 + 20, 400))

            self.tree.heading(col, text=col)
            self.tree.column(col, width=width, minwidth=80, anchor=tk.CENTER)

        for _, row in self.combined_df.iterrows():
            values = []
            for col in columns:
                val = row[col]
                if val is None or (isinstance(val, float) and pd.isna(val)):
                    values.append("")
                elif isinstance(val, (int, float, np.integer, np.floating)):
                    if col.startswith('AMIS '):
                        values.append(f"{val:.2f}")
                    else:
                        values.append(f"{val:g}")
                else:
                    values.append(str(val))
            self.tree.insert('', 'end', values=values)

    # =========================================================
    # HEADER CLICK → GRAPH
    # =========================================================

    def on_tree_click(self, event):
        region = self.tree.identify_region(event.x, event.y)
        if region != "heading":
            return
        col_id = self.tree.identify_column(event.x)
        if not col_id:
            return
        try:
            col_index = int(col_id.replace('#', '')) - 1
        except ValueError:
            return

        columns = list(self.tree['columns'])
        if col_index < 0 or col_index >= len(columns):
            return

        col_name = columns[col_index]
        if self.id_column_name and col_name == self.id_column_name:
            return

        base_name = col_name[5:] if col_name.startswith('AMIS ') else col_name

        matches = [(k, r) for k, r in self.results.items() if r['column'] == base_name]
        if not matches:
            matches = [(k, r) for k, r in self.results.items()
                       if base_name.startswith(r['column']) or r['column'].startswith(base_name)]

        if not matches:
            self.log(f"Parameter '{base_name}' not found for graph", "warning")
            return

        self.show_amis_graph(matches)

    def show_amis_graph(self, matches):
        fig = Figure(figsize=(9, 6), dpi=100)
        ax = fig.add_subplot(111)

        colors = plt.cm.tab10(np.linspace(0, 1, max(len(matches), 1)))
        title_param = None

        for (key, r), color in zip(matches, colors):
            raw = np.asarray(r['raw_values'], dtype=float)
            amis = np.asarray(r['amis_values'], dtype=float)

            ax.scatter(raw, amis, s=45, c=[color], edgecolors='black',
                       alpha=0.85, zorder=5, label=f"{r['column']}")

            order = np.argsort(raw)
            ax.plot(raw[order], amis[order], '-', color=color,
                    linewidth=1.8, alpha=0.7)

            if title_param is None:
                title_param = r['column']

        ax.set_xlabel(f'{title_param} (raw value)', fontweight='bold', fontsize=11)
        ax.set_ylabel('AMIS (0–100)', fontweight='bold', fontsize=11)
        ax.set_title(f'AMIS normalization: {title_param}', fontweight='bold', fontsize=12)
        ax.set_ylim(-5, 105)
        ax.grid(True, alpha=0.3)
        ax.axhline(50, color='orange', linestyle='--', alpha=0.6, label='AMIS = 50')
        for y in [0, 25, 50, 75, 100]:
            ax.axhline(y, color='green', linestyle=':', alpha=0.3, linewidth=1)
        ax.legend(fontsize=9, loc='best')
        fig.tight_layout()

        win = tk.Toplevel(self)
        win.title(f"AMIS graph — {title_param}")
        win.geometry("900x650")

        canvas = FigureCanvasTkAgg(fig, master=win)
        canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
        canvas.draw()

        btn_frame = ttk.Frame(win)
        btn_frame.pack(fill=tk.X, pady=5)

        def _save():
            default_dir = self.input_folder if self.input_folder else "."
            safe_name = str(title_param).replace(' ', '_').replace('/', '_')
            path = filedialog.asksaveasfilename(
                defaultextension=".png",
                initialfile=f"AMIS_graph_{safe_name}.png",
                initialdir=default_dir,
                filetypes=[("PNG files", "*.png"), ("PDF files", "*.pdf"), ("All files", "*.*")],
                title="Save graph"
            )
            if path:
                try:
                    fig.savefig(path, dpi=300, bbox_inches='tight')
                    messagebox.showinfo("Saved", f"Graph saved:\n{path}")
                except Exception as e:
                    messagebox.showerror("Error", f"Cannot save:\n{e}")

        ttk.Button(btn_frame, text="💾 Save PNG/PDF", command=_save).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="Close", command=win.destroy).pack(side=tk.RIGHT, padx=5)

        self.log(f"AMIS graph opened for '{title_param}'")

    # =========================================================
    # SAVE EXCEL
    # =========================================================

    def save_excel(self):
        if self.combined_df is None or self.combined_df.empty:
            messagebox.showwarning("No data", "Create table first")
            return

        default_dir = self.input_folder if self.input_folder else "."
        default_name = "AMIS_combined_batch_v37.xlsx"

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            initialfile=default_name,
            initialdir=default_dir,
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Save table as Excel"
        )

        if not file_path:
            return

        try:
            self.combined_df.to_excel(file_path, index=False)
            self._apply_excel_coloring(file_path)
            self.log(f"✅ Table saved with colors: {file_path}", "success")
            messagebox.showinfo("Saved", f"Table saved:\n{file_path}")

        except PermissionError:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            name, ext = os.path.splitext(file_path)
            file_path = f"{name}_{timestamp}{ext}"
            self.combined_df.to_excel(file_path, index=False)
            self._apply_excel_coloring(file_path)
            self.log(f"⚠️ File was locked, saved as: {os.path.basename(file_path)}", "warning")
            messagebox.showinfo("Saved", f"Table saved:\n{file_path}")

        except Exception as e:
            self.log(f"❌ Error saving: {str(e)}", "error")
            messagebox.showerror("Error", f"Cannot save:\n{e}")

    def _apply_excel_coloring(self, file_path):
        """Color pair columns in Excel: blue / white."""
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import PatternFill, Font, Alignment
        except ImportError:
            self.log("openpyxl not installed — coloring skipped", "warning")
            return

        wb = load_workbook(file_path)
        ws = wb.active

        header = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]

        col_to_pair = {}
        for j, name in enumerate(header, start=1):
            if name is None:
                col_to_pair[j] = None
                continue
            if self.id_column_name and name == self.id_column_name:
                col_to_pair[j] = None
                continue
            matched = False
            for pi, (p_raw, p_amis) in enumerate(self.column_pairs):
                if name == p_raw or name == p_amis:
                    col_to_pair[j] = pi
                    matched = True
                    break
            if not matched:
                col_to_pair[j] = None

        fill_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
        bold = Font(bold=True)
        center = Alignment(horizontal="center", vertical="center")

        for j in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=j)
            cell.font = bold
            cell.alignment = center
            pair_id = col_to_pair.get(j)
            if pair_id is None:
                continue
            if pair_id % 2 == 0:
                cell.fill = fill_blue
            else:
                cell.fill = fill_white

        for i in range(2, ws.max_row + 1):
            for j in range(1, ws.max_column + 1):
                pair_id = col_to_pair.get(j)
                if pair_id is None:
                    continue
                cell = ws.cell(row=i, column=j)
                cell.alignment = center
                if pair_id % 2 == 0:
                    cell.fill = fill_blue
                else:
                    cell.fill = fill_white

        for j in range(1, ws.max_column + 1):
            col_letter = ws.cell(row=1, column=j).column_letter
            max_len = 8
            for i in range(1, min(ws.max_row, 200) + 1):
                v = ws.cell(row=i, column=j).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            ws.column_dimensions[col_letter].width = min(max_len + 2, 22)

        ws.freeze_panes = "A2"

        wb.save(file_path)

    # =========================================================
    # COPY / RETURN
    # =========================================================

    def copy_to_clipboard(self):
        if self.combined_df is None or self.combined_df.empty:
            messagebox.showwarning("No data", "Create table first")
            return
        try:
            copy_str = self.combined_df.to_string(index=False)
            self.clipboard_clear()
            self.clipboard_append(copy_str)
            self.log("Table copied to clipboard", "success")
            messagebox.showinfo("Copied", "Table copied to clipboard")
        except Exception as e:
            self.log(f"Error copying: {str(e)}", "error")

    def return_to_main(self):
        try:
            import run_amis
            self.destroy()
            run_amis.show_main_menu()
        except Exception as e:
            self.log(f"Error returning to menu: {str(e)}", "error")
            self.destroy()


if __name__ == "__main__":
    app = BatchAMISApp()
    app.mainloop()